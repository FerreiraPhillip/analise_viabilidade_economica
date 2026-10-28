"""Importador genérico de planilhas de viabilidade com rótulos identificáveis."""

import re
import unicodedata
from collections import defaultdict

from openpyxl import load_workbook

from analysis import numero
from excel_utils import ler_planilha_tea
from project_analysis import analisar_projeto, converter_resultado_tea
from tea import analisar as analisar_tea


ALIASES = {
    "project": ("nome do projeto", "project name", "project title", "projeto", "project"),
    "investment": ("investimento inicial", "investimento total", "initial investment",
                   "initial investment required", "investment required", "total capital investment",
                   "total project investment", "initial capital requirement", "initial cost", "capital total", "capex"),
    "rate": ("taxa minima de atratividade", "taxa de desconto", "taxa minima",
             "discount rate", "minimum acceptable rate", "marr", "wacc", "hurdle rate"),
    "tax": ("taxa de imposto", "imposto sobre lucro", "tax rate", "income tax rate"),
    "horizon": ("vida util do projeto", "vida do projeto", "horizonte do projeto",
                "project life", "life-time project", "project horizon", "analysis period", "horizon"),
    "residual": ("valor residual", "salvage value", "residual value"),
    "benefit": ("receita anual", "receitas anuais", "beneficio anual", "beneficios anuais",
                "annual revenue", "annual revenues", "annual benefit", "annual benefits",
                "annual income", "annual savings"),
    "cost": ("custo anual", "custos anuais", "despesa anual", "despesas anuais",
             "annual cost", "annual costs", "operating cost", "operating costs", "operating expenses", "annual expenses"),
    "net_flow": ("fluxo de caixa liquido", "fluxo de caixa anual", "net cash flow",
                 "annual cash flow", "cash flow"),
}


def _norm(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _number(value):
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or "").strip().replace(" ", "")
    if not text:
        return None
    match = re.search(r"[-+]?\d[\d.,]*", text)
    if not match:
        return None
    return numero(match.group())


def _alias_key(text):
    normalized = _norm(text)
    # Prefer the most specific match where labels overlap.
    for key, aliases in ALIASES.items():
        if key == "project":
            continue
        for alias in sorted(aliases, key=len, reverse=True):
            if normalized == _norm(alias):
                return key
    for key, aliases in ALIASES.items():
        for alias in sorted(aliases, key=len, reverse=True):
            normalized_alias = _norm(alias)
            if len(normalized_alias) > 5 and normalized_alias in normalized:
                return key
    return None


def _has_any(text, phrases):
    return any(_norm(phrase) in text for phrase in phrases)


def _legacy_result(arquivo):
    """Converte os formatos já aceitos pelo site para a saída genérica."""
    arquivo.seek(0)
    globais, cenarios = ler_planilha_tea(arquivo)
    arquivo.seek(0)
    return converter_resultado_tea(analisar_tea(globais, cenarios))


def ler_planilha_generica(arquivo):
    """Lê a planilha do site, a planilha EVTE ou tabelas financeiras rotuladas."""
    try:
        return _legacy_result(arquivo)
    except (ValueError, KeyError, TypeError, AttributeError):
        arquivo.seek(0)

    wb = load_workbook(arquivo, data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        sheets.append((ws.title, rows))

    assumptions = {"project": "Projeto", "investment": None, "rate": None, "tax": 0,
                   "horizon": None, "residual": 0.0, "benefit": None, "cost": None}
    line_items = defaultdict(lambda: {"beneficios": [], "custos": []})
    flows_by_scenario = defaultdict(dict)
    inputs_by_scenario = defaultdict(dict)
    scenario_names = set()
    detected_rows = 0

    for title, rows in sheets:
        scenario_columns = {}
        for header_i, row in enumerate(rows[:15]):
            cells = list(row or ())
            marker = next((i for i, value in enumerate(cells)
                           if isinstance(value, str) and _has_any(_norm(value), ("cenario", "scenario"))), None)
            if marker is None:
                continue
            candidates = [cells]
            candidates.extend(list(x or ()) for x in rows[header_i + 1:header_i + 4])
            names_row = max(candidates, key=lambda candidate: sum(
                1 for value in candidate[marker:] if isinstance(value, str) and _norm(value)
                and not _has_any(_norm(value), ("cenario", "scenario"))))
            for col in range(marker + 1, len(names_row)):
                name = names_row[col]
                if isinstance(name, str) and _norm(name):
                    scenario_columns[col] = name.strip()
            if len(scenario_columns) >= 2:
                break

        for row in rows:
            cells = list(row or ())
            if not cells:
                continue
            # Lê pares rótulo/valor em qualquer aba, sem depender do nome da aba.
            for i, cell in enumerate(cells):
                if not isinstance(cell, str):
                    continue
                key = _alias_key(cell)
                if not key:
                    continue
                for col, scenario_name in scenario_columns.items():
                    if col < len(cells):
                        scenario_value = cells[col]
                        if key == "project":
                            if scenario_value is not None and str(scenario_value).strip():
                                inputs_by_scenario[scenario_name][key] = str(scenario_value).strip()
                        else:
                            parsed_scenario = _number(scenario_value)
                            if parsed_scenario is not None:
                                if key in ("rate", "tax") and 0 < parsed_scenario <= 1:
                                    parsed_scenario *= 100
                                inputs_by_scenario[scenario_name][key] = parsed_scenario
                if key == "project":
                    value = next((x for x in cells[i + 1:] if x is not None and str(x).strip()), None)
                else:
                    value = next((x for x in cells[i + 1:] if _number(x) is not None), None)
                if value is None:
                    continue
                if key == "project":
                    assumptions[key] = str(value).strip()
                    detected_rows += 1
                elif key in ("investment", "rate", "tax", "horizon", "residual", "benefit", "cost"):
                    parsed = _number(value)
                    if parsed is not None:
                        if key in ("rate", "tax") and 0 < parsed <= 1:
                            parsed *= 100
                        assumptions[key] = parsed
                        detected_rows += 1

        # Tabela longa: identifica colunas de ano + fluxo líquido ou categoria + valor.
        for header_i, row in enumerate(rows):
            headers = [_norm(x) for x in (row or ())]
            if not headers:
                continue
            col_year = next((i for i, h in enumerate(headers) if h in ("ano", "year", "periodo", "period") or _has_any(h, ("project year", "ano do projeto"))), None)
            col_flow = next((i for i, h in enumerate(headers) if _has_any(h, ("cash flow", "fluxo de caixa", "fluxo caixa"))), None)
            col_scenario = next((i for i, h in enumerate(headers) if _has_any(h, ("cenario", "scenario", "alternativa", "alternative"))), None)
            if col_year is not None and col_flow is not None:
                for data_row in rows[header_i + 1:]:
                    if not data_row or max(col_year, col_flow) >= len(data_row):
                        continue
                    year = _number(data_row[col_year])
                    amount = _number(data_row[col_flow])
                    if year is None or amount is None or year < 1:
                        continue
                    scenario = str(data_row[col_scenario]).strip() if col_scenario is not None and col_scenario < len(data_row) and data_row[col_scenario] else "Projeto"
                    flows_by_scenario[scenario][int(year)] = amount
                    scenario_names.add(scenario)

            col_kind = next((i for i, h in enumerate(headers) if _has_any(h, ("tipo", "type", "categoria", "category", "natureza"))), None)
            col_amount = next((i for i, h in enumerate(headers) if _has_any(h, ("valor", "value", "amount", "total anual", "annual amount"))), None)
            col_name = next((i for i, h in enumerate(headers) if _has_any(h, ("item", "descricao", "description", "nome", "name"))), None)
            if col_kind is not None and col_amount is not None:
                for data_row in rows[header_i + 1:]:
                    if max(col_kind, col_amount) >= len(data_row or ()):
                        continue
                    kind = _norm(data_row[col_kind])
                    amount = _number(data_row[col_amount])
                    if amount is None:
                        continue
                    name = str(data_row[col_name]).strip() if col_name is not None and col_name < len(data_row) and data_row[col_name] else (str(data_row[col_kind]).strip() or "Item")
                    bucket = "beneficios" if any(x in kind for x in ("receita", "beneficio", "revenue", "income", "benefit", "entrada")) else "custos" if any(x in kind for x in ("custo", "despesa", "cost", "expense", "saida")) else None
                    if bucket:
                        line_items["Projeto"][bucket].append({"nome": name, "valor": amount})
                        detected_rows += 1

    cenarios = []
    names = sorted(scenario_names or inputs_by_scenario.keys()) if (scenario_names or inputs_by_scenario) else [str(assumptions["project"] or "Projeto")]
    for name in names:
        scenario_inputs = inputs_by_scenario.get(name, {})
        flows = flows_by_scenario.get(name)
        benefits = list(line_items[name]["beneficios"] or line_items["Projeto"]["beneficios"])
        costs = list(line_items[name]["custos"] or line_items["Projeto"]["custos"])
        benefit = scenario_inputs.get("benefit", assumptions["benefit"])
        cost = scenario_inputs.get("cost", assumptions["cost"])
        if not benefits and benefit is not None:
            benefits = [{"nome": "Receitas e benefícios anuais", "valor": benefit}]
        if not costs and cost is not None:
            costs = [{"nome": "Custos e despesas anuais", "valor": cost}]
        if flows:
            last_year = max(flows)
            cash_flows = [flows.get(year, 0.0) for year in range(1, last_year + 1)]
        else:
            project_years = scenario_inputs.get("horizon", assumptions["horizon"])
            annual_net = sum(_number(item["valor"]) or 0 for item in benefits) - sum(_number(item["valor"]) or 0 for item in costs)
            cash_flows = [annual_net] * max(0, int(round(project_years))) if project_years else None
        cenarios.append({
            "nome": str(scenario_inputs.get("project", name)),
            "investimento_inicial": scenario_inputs.get("investment", assumptions["investment"]),
            "valor_residual": scenario_inputs.get("residual", assumptions["residual"]) or 0,
            "beneficios": benefits,
            "custos": costs,
            "fluxos_caixa": cash_flows,
        })

    missing = []
    if any(c["investimento_inicial"] is None for c in cenarios):
        missing.append("investimento inicial")
    if assumptions["rate"] is None and not all("rate" in x for x in inputs_by_scenario.values()):
        missing.append("taxa de desconto/MARR")
    if assumptions["horizon"] is None and not all("horizon" in x for x in inputs_by_scenario.values()) and not any(c["fluxos_caixa"] for c in cenarios):
        missing.append("vida útil do projeto ou anos de fluxo de caixa")
    has_annual_values = any(c["beneficios"] or c["custos"] for c in cenarios)
    has_flows = any(c["fluxos_caixa"] for c in cenarios)
    if not has_annual_values and not has_flows:
        missing.append("receitas/benefícios e custos anuais, ou uma tabela de fluxo de caixa líquido")
    if missing or detected_rows == 0:
        detail = ", ".join(missing) if missing else "dados reconhecíveis"
        raise ValueError(
            "Não consegui identificar todos os dados necessários nesta planilha: " + detail + ". "
            "Use rótulos como investimento inicial, taxa de desconto/MARR, vida útil, receitas/benefícios e custos anuais; "
            "ou uma tabela com Ano e Fluxo de caixa líquido."
        )

    horizon = assumptions["horizon"] or max(len(c["fluxos_caixa"] or []) for c in cenarios)
    rate = assumptions["rate"]
    if rate is None:
        rate = next(iter(inputs_by_scenario.values()))["rate"]
    return analisar_projeto({"taxa_desconto": rate, "vida_projeto": horizon,
                             "taxa_imposto": assumptions["tax"]}, cenarios)
