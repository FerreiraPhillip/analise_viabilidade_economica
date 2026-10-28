"""Gera static/modelo_planilha.xlsx: só valores digitados (de Teste.xlsx) + aba explicando o que o site calcula."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from campos import (GLOBAIS_GRUPOS, CENARIO_GRUPOS, EXEMPLO_GLOBAIS, EXEMPLO_CENARIOS)

CAB = PatternFill("solid", fgColor="13293D")
CAB_FONT = Font(bold=True, color="FFFFFF")
ENTRADA = PatternFill("solid", fgColor="FFF4CC")   # células a preencher
GRUPO = PatternFill("solid", fgColor="E3EAEE")


def cabecalho(ws, linha, titulos):
    for i, t in enumerate(titulos, start=1):
        c = ws.cell(row=linha, column=i, value=t)
        c.fill, c.font = CAB, CAB_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)


def larguras(ws, valores):
    for i, w in enumerate(valores, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ---------------- Leia-me ----------------
ws = wb.active
ws.title = "Leia-me"
linhas = [
    "Planilha-modelo — somente valores digitados",
    "",
    "• Células amarelas são as que você preenche. Nada nesta planilha é calculado: o site calcula sozinho tudo o que o Excel original calculava (veja a aba Calculado_pelo_site).",
    "• Aba Globais: parâmetros que valem para todos os cenários.",
    "• Aba Cenarios: uma coluna por cenário. Para adicionar um cenário, crie uma nova coluna à direita (o nome do cenário é o título da coluna).",
    "• Aba Equipamentos: uma linha por equipamento; a coluna 'cenario' deve repetir exatamente o nome do cenário.",
    "• Não altere os nomes das abas nem a coluna 'chave' (o site usa essas chaves para ler os valores).",
    "• Os valores vêm de Teste.xlsx e servem de exemplo: substitua pelos seus.",
    "• FCI, investimento inicial total (TCI) e edificações (aba Cenarios) são opcionais no modelo. Sem FCI, o site usa o custo de equipamentos como base de custos/depreciação; sem TCI, usa o FCI como investimento inicial e avisa que capital de giro/partida podem estar fora do VPL.",
]
for i, t in enumerate(linhas, start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws["A1"].font = Font(bold=True, size=14)
ws.column_dimensions["A"].width = 120

# ---------------- Globais ----------------
ws = wb.create_sheet("Globais")
cabecalho(ws, 1, ["chave", "variavel", "unidade", "valor"])
r = 2
for titulo, campos in GLOBAIS_GRUPOS:
    ws.cell(row=r, column=2, value=titulo).font = Font(bold=True)
    for col in range(1, 5):
        ws.cell(row=r, column=col).fill = GRUPO
    r += 1
    for chave, rotulo, unidade in campos:
        ws.cell(row=r, column=1, value=chave)
        ws.cell(row=r, column=2, value=rotulo)
        ws.cell(row=r, column=3, value=unidade)
        c = ws.cell(row=r, column=4, value=EXEMPLO_GLOBAIS.get(chave))
        c.fill = ENTRADA
        r += 1
larguras(ws, [22, 52, 40, 14])
ws.freeze_panes = "A2"

# ---------------- Cenarios ----------------
ws = wb.create_sheet("Cenarios")
nomes = [c["nome"] for c in EXEMPLO_CENARIOS]
cabecalho(ws, 1, ["chave", "variavel", "unidade"] + nomes)
r = 2
for titulo, campos in CENARIO_GRUPOS:
    ws.cell(row=r, column=2, value=titulo).font = Font(bold=True)
    for col in range(1, 4 + len(nomes)):
        ws.cell(row=r, column=col).fill = GRUPO
    r += 1
    for chave, rotulo, unidade in campos:
        ws.cell(row=r, column=1, value=chave)
        ws.cell(row=r, column=2, value=rotulo)
        ws.cell(row=r, column=3, value=unidade)
        for j, cen in enumerate(EXEMPLO_CENARIOS):
            valor = cen.get(chave, "")
            c = ws.cell(row=r, column=4 + j, value=(valor if valor != "" else None))
            c.fill = ENTRADA
        r += 1
larguras(ws, [18, 48, 14] + [34] * len(nomes))
ws.freeze_panes = "D2"

# ---------------- Equipamentos ----------------
ws = wb.create_sheet("Equipamentos")
cabecalho(ws, 1, ["cenario", "equipamento", "quantidade", "custo_unidade_usd", "capacidade",
                  "unidade_capacidade", "consumo_kwh_ano", "material", "fonte", "ajustar_cepci"])
r = 2
for cen in EXEMPLO_CENARIOS:
    for e in cen["equipamentos"]:
        vals = [cen["nome"], e["nome"], e["quantidade"], e["custo_unidade"], e["capacidade"],
                e["unidade_capacidade"], (e["consumo_kwh_ano"] if e["consumo_kwh_ano"] != "" else None),
                e["material"], e["fonte"], e["ajustar_cepci"]]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=col, value=v)
            if col > 1:
                c.fill = ENTRADA
        r += 1
larguras(ws, [34, 36, 11, 18, 12, 18, 18, 16, 36, 15])
ws.freeze_panes = "A2"

# ---------------- Calculado_pelo_site ----------------
ws = wb.create_sheet("Calculado_pelo_site")
cabecalho(ws, 1, ["Item calculado", "Onde estava em Teste.xlsx", "Como o site calcula"])
itens = [
    ("Razão CEPCI", "Custo dos Equipamentos!K6", "índice CEPCI atual ÷ índice CEPCI do ano-base"),
    ("Valor ajustado do equipamento", "Custo dos Equipamentos!E13:E23 (=D ou =D×K6)", "custo unitário; multiplicado pela razão CEPCI quando 'ajustar_cepci' = sim"),
    ("Custo total do equipamento", "Custo dos Equipamentos!F13:F23", "valor ajustado × quantidade"),
    ("Custo total de equipamentos do cenário", "Custo dos Equipamentos!D19/F19 e D24/F24", "soma dos custos totais dos equipamentos do cenário"),
    ("Consumo elétrico do cenário", "Custo dos Equipamentos!I19, I24 e Tecno Assumptions!D23:E23", "soma do consumo elétrico (kWh/ano) dos equipamentos do cenário"),
    ("Rendimento do processo 1", "Tecno Assumptions!D17:E17", "produto 1 ÷ matéria-prima processada"),
    ("Rendimento do processo 2", "Tecno Assumptions!D18:E18", "produto 2 ÷ matéria-prima processada"),
    ("Preço final do etanol e do CO₂", "Tecno Assumptions!D33:D34 (=500-500*10%, =100-10%*100)", "preço de lista × (1 − desconto). Preço e desconto agora são dois campos digitados"),
    ("Preço dos produtos em US$/t", "Tecno Assumptions!D40:E41 (=250*1000 etc.)", "preço em US$/kg × 1.000. Você digita o preço em US$/kg"),
    ("", "", ""),
    ("NÃO replicado: total geral de equipamentos e consumo", "Custo dos Equipamentos!F39 e I39", "No Excel, F39 = SUM(F13:F38) soma os itens E os subtotais (F19, F24), contando tudo em dobro (256.313 em vez de 128.157); além disso, soma cenários que são alternativas entre si. O site mostra o total por cenário."),
    ("Acrescentado pelo site (não estava no Excel)", "—", "receita, custos anuais pelos critérios percentuais, depreciação linear, imposto, fluxo de caixa, VP, VF, VPL, TIR, payback e veredito de viabilidade"),
]
for i, linha in enumerate(itens, start=2):
    for j, v in enumerate(linha, start=1):
        ws.cell(row=i, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
larguras(ws, [42, 52, 90])

wb.save("static/modelo_planilha.xlsx")
print("modelo gerado")
